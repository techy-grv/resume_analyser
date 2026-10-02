from pathlib import Path
from typing import BinaryIO

from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_pdf_text(file_path: Path | BinaryIO) -> str:
    # PdfReader accepts both a path and a file-like object (e.g. a Streamlit upload).
    reader = PdfReader(file_path)
    pages = []

    for page in reader.pages:
        pages.append(page.extract_text() or "")

    return "\n".join(pages).strip()


def extract_docx_text(file_path: Path) -> str:
    document = DocxDocument(str(file_path))

    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(paragraphs).strip()


def load_resumes(resume_dir: Path) -> list[Document]:
    documents = []

    files = sorted(
        file_path
        for file_path in resume_dir.iterdir()
        if file_path.is_file()
        and file_path.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    for file_path in files:
        try:
            if file_path.suffix.lower() == ".pdf":
                text = extract_pdf_text(file_path)
            else:
                text = extract_docx_text(file_path)

            if not text:
                continue

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "candidate_name": file_path.stem.replace("_", " "),
                    },
                )
            )
        except Exception as exc:
            print(f"Skipping {file_path.name}: {exc}")

    return documents
